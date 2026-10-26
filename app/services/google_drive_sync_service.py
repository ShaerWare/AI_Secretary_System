"""Google Drive sync service — lists a Drive source, downloads files, converts to markdown.

A source is either a whole folder (recursive) or an explicit list of file ids.
Two consumers:
- RAG: every supported file becomes a markdown document in the collection dir
  (`convert_files_to_markdown`). PDF/XLSX/XLS/DOCX are converted here — the
  supplier documents people share are binary, and the old text-only import
  skipped all of them as "unsupported extension".
- Supplier prices: the raw file is needed for the price parser
  (`download_raw_file`); a native Google Sheet is exported to xlsx.

`files_fingerprint` + `state_changed` let the daily check skip the work when
nothing on Drive moved since the last sync.
"""

import asyncio
import io
import logging
import re
import unicodedata
from pathlib import Path

import httpx


logger = logging.getLogger(__name__)

DRIVE_FILES_URL = "https://www.googleapis.com/drive/v3/files"
FOLDER_MIME = "application/vnd.google-apps.folder"
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILE_FIELDS = "id, name, mimeType, size, modifiedTime"

# Папки, расшаренные другим аккаунтом, и общие диски видны только с этими флагами.
ALL_DRIVES = {"supportsAllDrives": "true", "includeItemsFromAllDrives": "true"}

# Google Apps MIME types that can be exported as text (RAG)
EXPORTABLE_TYPES = {
    "application/vnd.google-apps.document": ("text/plain", ".txt", "google_docs"),
    SHEET_MIME: (
        "text/csv",
        ".csv",
        "google_sheets",
    ),
    "application/vnd.google-apps.presentation": (
        "text/plain",
        ".txt",
        "google_slides",
    ),
}

# Regular file types we can read as text
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".yml",
    ".yaml",
    ".toml",
    ".ini",
    ".cfg",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".vue",
    ".jsx",
    ".tsx",
    ".css",
    ".scss",
    ".sql",
    ".sh",
    ".bash",
    ".rb",
    ".go",
    ".rs",
    ".java",
    ".kt",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
}

# Binary documents converted to markdown for RAG
DOCUMENT_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".docx"}

# What the supplier price parser can read (see modules/procurement/suppliers/registry.py)
PRICE_EXTENSIONS = {".xlsx", ".xls", ".pdf"}

MAX_FILE_SIZE = 512 * 1024  # 512 KB per plain-text file
MAX_DOCUMENT_SIZE = 50 * 1024 * 1024  # 50 MB per binary document / price list

MAX_SHEET_ROWS = 5000  # per sheet, RAG conversion
ROWS_PER_SECTION = 200  # one RAG section per N table rows, header repeated
MAX_PDF_PAGES = 200


def drive_name(name: str) -> str:
    """Drive names from macOS come decomposed (NFD): «й» = «и» + breve.

    The supplier registry matches file names by pattern («Диллер*.xlsx»), and an
    NFD «й» never equals the NFC one typed in the pattern.
    """
    return unicodedata.normalize("NFC", name).strip()


def extract_drive_id(link_or_id: str) -> str | None:
    """Folder/file id from a Drive link (…/folders/<id>, …/d/<id>, ?id=<id>) or a bare id."""
    s = link_or_id.strip()
    for pattern in (r"/folders/([\w-]{10,})", r"/d/([\w-]{10,})", r"[?&]id=([\w-]{10,})"):
        m = re.search(pattern, s)
        if m:
            return m.group(1)
    if re.fullmatch(r"[\w-]{10,}", s):
        return s
    return None


async def list_drive_folder_recursive(
    access_token: str,
    folder_id: str = "root",
    max_files: int = 500,
) -> list[dict]:
    """List all files in a Drive folder recursively (up to max_files)."""
    files: list[dict] = []
    folders_to_scan = [folder_id]

    async with httpx.AsyncClient(timeout=30) as client:
        while folders_to_scan and len(files) < max_files:
            current_folder = folders_to_scan.pop(0)
            page_token = None

            while True:
                params: dict = {
                    "q": f"'{current_folder}' in parents and trashed = false",
                    "fields": f"nextPageToken, files({FILE_FIELDS})",
                    "pageSize": 100,
                    **ALL_DRIVES,
                }
                if page_token:
                    params["pageToken"] = page_token

                resp = await client.get(
                    DRIVE_FILES_URL,
                    headers={"Authorization": f"Bearer {access_token}"},
                    params=params,
                )
                resp.raise_for_status()
                data = resp.json()

                for f in data.get("files", []):
                    if f["mimeType"] == FOLDER_MIME:
                        folders_to_scan.append(f["id"])
                    else:
                        files.append(f)
                        if len(files) >= max_files:
                            break

                page_token = data.get("nextPageToken")
                if not page_token or len(files) >= max_files:
                    break

    return files


async def get_files_meta(access_token: str, file_ids: list[str]) -> tuple[list[dict], list[str]]:
    """Metadata of explicitly selected files. Returns (files, missing_ids).

    A selected file that was deleted or un-shared comes back as missing instead
    of failing the whole source.
    """
    files: list[dict] = []
    missing: list[str] = []
    async with httpx.AsyncClient(timeout=30) as client:
        for file_id in file_ids:
            resp = await client.get(
                f"{DRIVE_FILES_URL}/{file_id}",
                headers={"Authorization": f"Bearer {access_token}"},
                params={"fields": f"{FILE_FIELDS}, trashed", "supportsAllDrives": "true"},
            )
            if resp.status_code in (403, 404):
                missing.append(file_id)
                continue
            resp.raise_for_status()
            meta = resp.json()
            if meta.get("trashed"):
                missing.append(file_id)
            elif meta["mimeType"] == FOLDER_MIME:
                files.extend(await list_drive_folder_recursive(access_token, meta["id"]))
            else:
                files.append(meta)
    return files, missing


async def list_source_files(
    access_token: str, folder_id: str, file_ids: list[str] | None
) -> tuple[list[dict], list[str]]:
    """Files of a source: the selected ones, or the whole folder. Returns (files, missing_ids)."""
    if file_ids:
        return await get_files_meta(access_token, file_ids)
    return await list_drive_folder_recursive(access_token, folder_id), []


def files_fingerprint(files: list[dict]) -> dict:
    """{file_id: {name, modifiedTime}} — what the source looked like at sync time."""
    return {
        f["id"]: {"name": drive_name(f["name"]), "modifiedTime": f.get("modifiedTime")}
        for f in files
    }


def state_changed(old_state: dict | None, new_state: dict) -> bool:
    """True if a file was added, removed, renamed or modified since `old_state`."""
    if not old_state:
        return True
    strip = {
        fid: {"name": v.get("name"), "modifiedTime": v.get("modifiedTime")}
        for fid, v in old_state.items()
    }
    return strip != new_state


# ── Download ─────────────────────────────────────────────────────


async def _download(client: httpx.AsyncClient, access_token: str, file_info: dict) -> bytes | None:
    resp = await client.get(
        f"{DRIVE_FILES_URL}/{file_info['id']}",
        headers={"Authorization": f"Bearer {access_token}"},
        params={"alt": "media", "supportsAllDrives": "true"},
    )
    if resp.status_code != 200:
        logger.warning(f"Failed to download {file_info['name']}: {resp.status_code}")
        return None
    return resp.content


async def _export(
    client: httpx.AsyncClient, access_token: str, file_info: dict, mime: str
) -> bytes | None:
    resp = await client.get(
        f"{DRIVE_FILES_URL}/{file_info['id']}/export",
        headers={"Authorization": f"Bearer {access_token}"},
        params={"mimeType": mime},
    )
    if resp.status_code != 200:
        logger.warning(f"Failed to export {file_info['name']}: {resp.status_code}")
        return None
    return resp.content


def price_file_name(file_info: dict) -> str | None:
    """Local name for a price file, or None if the parser can't read this type."""
    name = drive_name(file_info["name"]).replace("/", "_").replace("\\", "_").lstrip(".")
    if file_info["mimeType"] == SHEET_MIME:
        return name if name.lower().endswith(".xlsx") else f"{name}.xlsx"
    if Path(name).suffix.lower() in PRICE_EXTENSIONS:
        return name
    return None


async def download_raw_file(access_token: str, file_info: dict) -> bytes | None:
    """Raw bytes for the price parser; a native Google Sheet is exported to xlsx."""
    size = int(file_info.get("size") or 0)
    if size > MAX_DOCUMENT_SIZE:
        raise ValueError(f"{file_info['name']}: слишком большой ({size / 1048576:.0f} МБ)")
    async with httpx.AsyncClient(timeout=120) as client:
        if file_info["mimeType"] == SHEET_MIME:
            return await _export(client, access_token, file_info, XLSX_MIME)
        return await _download(client, access_token, file_info)


# ── Conversion to markdown (RAG) ────────────────────────────────


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("|", "/").replace("\n", " ").strip()


def _rows_to_sections(sheet_title: str, rows: list[list]) -> list[str]:
    """A sheet as markdown tables of ROWS_PER_SECTION rows, header repeated in each.

    RAG splits documents by headings; one 5000-row table would be a single
    chunk the retriever can only take or drop whole.
    """
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    head = ["| " + " | ".join(header) + " |", "|" + " --- |" * width]
    parts = [f"## {sheet_title}"]
    if not body:
        parts.append("\n".join(head))
    for start in range(0, len(body), ROWS_PER_SECTION):
        chunk = body[start : start + ROWS_PER_SECTION]
        lines = head + ["| " + " | ".join(r) + " |" for r in chunk]
        parts.append(f"### {sheet_title}: строки {start + 1}–{start + len(chunk)}")
        parts.append("\n".join(lines))
    return parts


def _xlsx_to_markdown(data: bytes) -> str | None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True)]
            parts.extend(_rows_to_sections(ws.title, rows))
        return "\n\n".join(parts) or None
    finally:
        wb.close()


def _xls_to_markdown(data: bytes) -> str | None:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    parts: list[str] = []
    for sheet in book.sheets():
        rows = [sheet.row_values(i) for i in range(min(sheet.nrows, MAX_SHEET_ROWS))]
        parts.extend(_rows_to_sections(sheet.name, rows))
    return "\n\n".join(parts) or None


def _pdf_to_markdown(data: bytes) -> str | None:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages[:MAX_PDF_PAGES], start=1):
            text = page.extract_text()
            if text and text.strip():
                pages.append(f"## Страница {i}\n\n{text.strip()}")
    return "\n\n".join(pages) or None


def _docx_to_markdown(data: bytes) -> str | None:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for n, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parts.extend(_rows_to_sections(f"Таблица {n}", rows))
    return "\n\n".join(parts) or None


def _text_to_markdown(data: bytes) -> str | None:
    for enc in ("utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            return text if text.strip() else None
        except UnicodeDecodeError:
            continue
    return None


_CONVERTERS = {
    ".xlsx": _xlsx_to_markdown,
    ".xls": _xls_to_markdown,
    ".pdf": _pdf_to_markdown,
    ".docx": _docx_to_markdown,
}


def document_to_markdown(name: str, data: bytes) -> str | None:
    """Convert a downloaded file to markdown. None if unreadable/empty. CPU-bound."""
    ext = Path(name).suffix.lower()
    converter = _CONVERTERS.get(ext)
    try:
        if converter:
            return converter(data)
        return _text_to_markdown(data)
    except Exception as e:
        logger.warning(f"Conversion failed for {name}: {e}")
        return None


async def download_and_convert_file(
    access_token: str,
    file_info: dict,
) -> tuple[str, str, int] | None:
    """Download a single file and convert to text.

    Returns (title, content, size_bytes) or None if file can't be processed.
    """
    name = drive_name(file_info["name"])
    mime_type = file_info["mimeType"]
    size = int(file_info.get("size") or 0)
    ext = Path(name).suffix.lower()

    async with httpx.AsyncClient(timeout=120) as client:
        # Google Apps files (Docs, Sheets, Slides) — export
        if mime_type in EXPORTABLE_TYPES:
            export_mime, _ext, _source = EXPORTABLE_TYPES[mime_type]
            if mime_type == SHEET_MIME:
                # xlsx keeps every sheet; the csv export returns only the first one
                data = await _export(client, access_token, file_info, XLSX_MIME)
                if data is None:
                    return None
                content = await asyncio.to_thread(_xlsx_to_markdown, data)
            else:
                data = await _export(client, access_token, file_info, export_mime)
                content = data.decode("utf-8", errors="replace") if data else None
            if not content:
                return None
            return (name, content, len(data or b""))

        if ext in DOCUMENT_EXTENSIONS:
            limit = MAX_DOCUMENT_SIZE
        elif ext in TEXT_EXTENSIONS:
            limit = MAX_FILE_SIZE
        else:
            logger.info(f"Skipping {name}: unsupported extension {ext}")
            return None
        if size > limit:
            logger.info(f"Skipping {name}: too large ({size} bytes)")
            return None

        data = await _download(client, access_token, file_info)
        if data is None:
            return None

    content = await asyncio.to_thread(document_to_markdown, name, data)
    if not content:
        return None
    return (name, content, len(data))


def _sanitize_filename(name: str) -> str:
    """Convert file name to safe filename for disk."""
    name = re.sub(r"[^\w\s\-.]", "_", name)
    name = re.sub(r"\s+", "_", name)
    return name[:200]


async def convert_files_to_markdown(
    access_token: str,
    drive_files: list[dict],
    output_dir: str,
) -> tuple[list[dict], list[dict]]:
    """Write every convertible file to `output_dir` as markdown.

    Returns (documents for the DatasetSynced event, skipped [{name, reason}]).
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    documents: list[dict] = []
    contents: dict[str, str] = {}
    skipped: list[dict] = []
    used_names: set[str] = set()

    for file_info in drive_files:
        try:
            result = await download_and_convert_file(access_token, file_info)
        except Exception as e:
            logger.warning(f"Error processing {file_info['name']}: {e}")
            skipped.append({"name": file_info["name"], "reason": str(e)[:200]})
            continue
        if not result:
            skipped.append({"name": file_info["name"], "reason": "формат не читается или пусто"})
            continue
        title, content, size_bytes = result

        # Determine source type
        mime_type = file_info["mimeType"]
        if mime_type in EXPORTABLE_TYPES:
            source_type = EXPORTABLE_TYPES[mime_type][2]
        else:
            source_type = "google_drive"

        # Write to disk as markdown; same-named files from different subfolders
        # must not overwrite each other
        stem = _sanitize_filename(title)
        stem = stem.rsplit(".", 1)[0] if "." in stem else stem
        safe_name = f"{stem}.md"
        n = 2
        while safe_name in used_names:
            safe_name = f"{stem}_{n}.md"
            n += 1
        used_names.add(safe_name)

        # Wrap content in markdown with title
        md_content = f"# {title}\n\n{content}"
        contents[safe_name] = md_content

        # Count sections
        section_count = len(re.findall(r"^#{2,3}\s+.+$", md_content, re.MULTILINE))

        documents.append(
            {
                "filename": safe_name,
                "title": title,
                "source_type": source_type,
                "file_size_bytes": size_bytes,
                "section_count": max(section_count, 1),
            }
        )

    # Old files go only once there is something to replace them with: a run
    # where every download failed must leave the collection as it was
    if documents:
        for old_file in Path(output_dir).glob("*.md"):
            old_file.unlink()
        for safe_name, md_content in contents.items():
            (Path(output_dir) / safe_name).write_text(md_content, encoding="utf-8")

    total_size = sum(d["file_size_bytes"] for d in documents)
    logger.info(
        f"Google Drive sync complete: {len(documents)} documents, {total_size / 1024:.1f} KB total"
    )
    return documents, skipped
